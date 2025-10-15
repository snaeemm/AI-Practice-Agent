import os
import sys
import json
from pathlib import Path
from typing import Dict, Any, Optional, List
from datetime import datetime
from dotenv import load_dotenv

load_dotenv()

from agent.database.db_singleton import get_db
from agent.database.db_manager import DatabaseManager # Added for fetching data
import io # Added for byte stream handling

db = get_db()


def tool_qualify_rfp(context: Optional[str] = None, pdf_path: Optional[str] = None) -> Dict[str, Any]:
    """
    Qualify an RFP using the qualification matrix and capabilities.

    Priority:
    1. Context (preferred): Pre-extracted document text from UI or user-provided text
    2. PDF path (legacy): For existing FILES_DIR PDFs only

    Args:
        context: Pre-extracted RFP text content (preferred method)
        pdf_path: Legacy parameter for FILES_DIR PDFs (fallback only)

    Returns:
        Dictionary with qualification results including score and decision
    """
    try:
        from pathlib import Path
        from .processors.new_rfp_qualifier import process_rfp_qualification, UserContext, RFPData

        # Validate that at least one input is provided
        if not context and not pdf_path:
            return {
                'success': False,
                'error': 'Missing input',
                'message': 'Either context or pdf_path must be provided'
            }

        # Prioritize context over pdf_path
        user_context = None
        if context:
            # Context is the preferred method - treat as full RFP content
            user_context = UserContext(rfp_content=context)
            pdf_path = None  # Ignore pdf_path if context is provided

        # STEP 1: Quick extraction to get client/project names for duplicate check
        print("🔍 Checking if RFP already qualified...")
        from .processors.new_rfp_qualifier import extract_rfp_data
        rfp_data_check, _, _ = extract_rfp_data(None, user_context)

        if rfp_data_check and rfp_data_check.client_and_opportunity:
            client_name = rfp_data_check.client_and_opportunity.value if rfp_data_check.client_and_opportunity.value else "Unknown"
            project_title = client_name  # Use same for lookup

            # Check if already exists
            existing_rfp_id = db.find_existing_rfp(client_name, project_title)
            if existing_rfp_id:
                # Check if already qualified
                existing_qual = db.get_qualification_results(existing_rfp_id)
                if existing_qual:
                    print(f"✅ RFP already qualified: {existing_rfp_id}")
                    print("⚠️  Skipping re-qualification - returning existing results")

                    # Return existing qualification data
                    report_data = existing_qual.get('qualification_report', {})
                    return {
                        'success': True,
                        'message': f'RFP already qualified (using existing): {existing_rfp_id}',
                        'rfp_id': existing_rfp_id,
                        'already_processed': True,
                        'total_score': report_data.get('total_score'),
                        'threshold': report_data.get('threshold'),
                        'qualifies': report_data.get('qualifies'),
                        'decision': 'PURSUE' if report_data.get('qualifies') else 'DECLINE',
                        'executive_summary': report_data.get('executive_summary'),
                        'recommendations': report_data.get('recommendations', []),
                        'rfp_classification': report_data.get('rfp_classification'),
                        'analyses': report_data.get('analyses', []),
                        'qualification_context': report_data.get('qualification_context')
                    }

        # STEP 2: Not qualified yet, proceed with qualification
        print("🚀 Proceeding with new qualification...")

        # Determine pdf_base for file naming
        if pdf_path:
            pdf_file = Path(pdf_path)
            pdf_base = pdf_file.stem
        else:
            # Generate a name for context-only processing
            import hashlib
            import datetime
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            context_hash = hashlib.md5(context[:100].encode()).hexdigest()[:8]
            pdf_base = f"context_rfp_{timestamp}_{context_hash}"

        report, rfp_id = process_rfp_qualification(pdf_path, user_context=user_context)

        if not report:
            return {
                'success': False,
                'error': 'Qualification processing failed',
                'message': 'Failed to generate qualification report'
            }

        return {
            'success': True,
            'message': f'RFP qualification completed for {rfp_id}',
            'rfp_id': rfp_id, # Changed from pdf_base to rfp_id
            'total_score': report.total_score,
            'threshold': report.threshold,
            'qualifies': report.qualifies,
            'decision': 'PURSUE' if report.qualifies else 'DECLINE',
            'executive_summary': report.executive_summary,
            'recommendations': report.recommendations,
            'rfp_classification': report.rfp_classification,
            'analyses': [
                {
                    'criterion': a.criterion,
                    'selected_option': a.selected_option,
                    'score': a.score,
                    'weighted_score': a.weighted_score,
                    'reasoning': a.reasoning,
                    'missing_data_justification': a.missing_data_justification
                }
                for a in report.analyses
            ],
            'qualification_context': {
                'estimated_budget': {
                    'value': report.qualification_context.estimated_budget.value if report.qualification_context and report.qualification_context.estimated_budget else None,
                    'source': report.qualification_context.estimated_budget.source if report.qualification_context and report.qualification_context.estimated_budget else None,
                    'confidence_level': report.qualification_context.estimated_budget.confidence_level if report.qualification_context and report.qualification_context.estimated_budget else None
                },
                'estimated_timeline': {
                    'value': report.qualification_context.estimated_timeline.value if report.qualification_context and report.qualification_context.estimated_timeline else None,
                    'source': report.qualification_context.estimated_timeline.source if report.qualification_context and report.qualification_context.estimated_timeline else None,
                    'confidence_level': report.qualification_context.estimated_timeline.confidence_level if report.qualification_context and report.qualification_context.estimated_timeline else None
                },
                'capability_gaps': report.qualification_context.capability_gaps if report.qualification_context else [],
                'competitive_advantages': report.qualification_context.competitive_advantages if report.qualification_context else [],
                'risk_factors': report.qualification_context.risk_factors if report.qualification_context else [],
                'opportunity_factors': report.qualification_context.opportunity_factors if report.qualification_context else []
            } if report.qualification_context else None
        }

    except Exception as e:
        return {
            'success': False,
            'error': str(e),
            'message': f'Failed to qualify RFP: {str(e)}'
        }


def tool_plan_bid_sections(context: Optional[str] = None, pdf_path: Optional[str] = None, rfp_id: Optional[str] = None) -> Dict[str, Any]:
    """
    Create detailed bid plan with section assignments and recommendations.

    Priority:
    1. Context (preferred): Pre-extracted document text from UI or user-provided text
    2. PDF path (legacy): For existing FILES_DIR PDFs only

    Args:
        context: Pre-extracted RFP text content (preferred method)
        pdf_path: Legacy parameter for FILES_DIR PDFs (fallback only)
        rfp_id: Optional RFP ID from qualification step (to ensure same ID is used)

    Returns:
        Dictionary with bid plan including sections and assignments
    """
    try:
        from pathlib import Path
        from .processors.new_rfp_bid_planner import process_pdf
        from .processors.new_rfp_qualifier import UserContext

        # Validate that at least one input is provided
        if not context and not pdf_path:
            return {
                'success': False,
                'error': 'Missing input',
                'message': 'Either context or pdf_path must be provided'
            }

        # Prioritize context over pdf_path
        user_context = None
        if context:
            # Context is the preferred method - treat as full RFP content
            user_context = UserContext(rfp_content=context)
            pdf_path = None  # Ignore pdf_path if context is provided

        # STEP 1: Check if bid plan already exists for this RFP
        if rfp_id:
            print(f"🔍 Checking if bid plan already exists for: {rfp_id}")
            existing_deliverables = db.get_rfp_deliverables(rfp_id)
            existing_assignments = db.get_rfp_assignments(rfp_id)

            if existing_deliverables and existing_assignments:
                print(f"✅ Bid plan already exists for: {rfp_id}")
                print("⚠️  Skipping re-planning - returning existing results")

                # Return existing bid plan data
                return {
                    'success': True,
                    'message': f'Bid plan already exists (using existing): {rfp_id}',
                    'rfp_id': rfp_id,
                    'already_processed': True,
                    'client_and_opportunity': existing_deliverables.get('client_and_opportunity'),
                    'total_deliverables': existing_assignments.get('total_deliverables'),
                    'granite_assigned': existing_assignments.get('granite_assigned'),
                    'partner_assigned': existing_assignments.get('partner_assigned'),
                    'assignments': existing_assignments.get('assignment_report', {}).get('assignments', [])
                }

        # STEP 2: Not yet planned, or no rfp_id provided - check by client/project names
        if not rfp_id and context:
            print("🔍 Checking if RFP already has bid plan (by client/project)...")
            from .processors.new_rfp_qualifier import extract_rfp_data
            rfp_data_check, _, _ = extract_rfp_data(None, user_context)

            if rfp_data_check and rfp_data_check.client_and_opportunity:
                client_name = rfp_data_check.client_and_opportunity.value if rfp_data_check.client_and_opportunity.value else "Unknown"
                project_title = client_name

                # Check if already exists
                existing_rfp_id = db.find_existing_rfp(client_name, project_title)
                if existing_rfp_id:
                    # Check if already has bid plan
                    existing_deliverables = db.get_rfp_deliverables(existing_rfp_id)
                    existing_assignments = db.get_rfp_assignments(existing_rfp_id)

                    if existing_deliverables and existing_assignments:
                        print(f"✅ Bid plan already exists: {existing_rfp_id}")
                        print("⚠️  Skipping re-planning - returning existing results")

                        return {
                            'success': True,
                            'message': f'Bid plan already exists (using existing): {existing_rfp_id}',
                            'rfp_id': existing_rfp_id,
                            'already_processed': True,
                            'client_and_opportunity': existing_deliverables.get('client_and_opportunity'),
                            'total_deliverables': existing_assignments.get('total_deliverables'),
                            'granite_assigned': existing_assignments.get('granite_assigned'),
                            'partner_assigned': existing_assignments.get('partner_assigned'),
                            'assignments': existing_assignments.get('assignment_report', {}).get('assignments', [])
                        }
                    else:
                        # Exists but no bid plan yet - use the existing rfp_id
                        rfp_id = existing_rfp_id
                        print(f"🔄 Using existing rfp_id (no bid plan yet): {rfp_id}")

        # STEP 3: Proceed with bid planning
        print("🚀 Proceeding with new bid planning...")

        # Determine pdf filename - IMPORTANT: Pass rfp_id as pdf_input to process_pdf
        if pdf_path:
            pdf_file = Path(pdf_path)
            pdf_filename = pdf_file.name
            pdf_name = pdf_file.stem
            pdf_input_for_processor = pdf_path
        elif rfp_id:
            # Use the provided rfp_id from qualification step
            # Pass rfp_id as the pdf_input so process_pdf uses it as pdf_base
            pdf_filename = f"{rfp_id}.txt"
            pdf_name = rfp_id
            pdf_input_for_processor = rfp_id  # Pass the rfp_id directly
            print(f"🔄 Using rfp_id: {rfp_id}")
        else:
            # Generate a name for context-only processing (fallback)
            import hashlib
            import datetime
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            context_hash = hashlib.md5(context[:100].encode()).hexdigest()[:8]
            pdf_filename = f"context_rfp_{timestamp}_{context_hash}.txt"
            pdf_name = f"context_rfp_{timestamp}_{context_hash}"
            pdf_input_for_processor = None  # Let process_pdf generate from context

        template_input = "Bid Plan - [Client Opp Name]_BB_140125.xlsx"
        output_template = "{pdf_name}_bid_plan.xlsx"

        # CRITICAL: Pass rfp_id as pdf_input so process_pdf uses the same ID
        process_pdf(pdf_input_for_processor, template_input, output_template, user_context)

        # Read the generated assignment analysis JSON to return to agent
        import os
        results_dir = Path(os.getenv("RESULTS_DIR", "./results"))
        assignment_json_path = results_dir / f"{pdf_name}_assignment_analysis.json"

        assignments_data = None
        if assignment_json_path.exists():
            try:
                import json
                with open(assignment_json_path, 'r', encoding='utf-8') as f:
                    assignments_data = json.load(f)
            except Exception as e:
                print(f"⚠️ Failed to read assignment JSON: {e}")

        result = {
            'success': True,
            'message': f'Bid plan created for {pdf_filename}',
            'pdf_name': pdf_name
        }

        # Add assignment data if available
        if assignments_data:
            result['client_and_opportunity'] = assignments_data.get('client_and_opportunity')
            result['total_deliverables'] = assignments_data.get('total_deliverables')
            result['granite_assigned'] = assignments_data.get('granite_assigned')
            result['partner_assigned'] = assignments_data.get('partner_assigned')
            result['assignments'] = assignments_data.get('assignments', [])

        return result

    except Exception as e:
        return {
            'success': False,
            'error': str(e),
            'message': f'Failed to plan bid: {str(e)}'
        }


# ==================== Report Generation Tools ====================

def generate_and_download_qualification_excel(rfp_id: str) -> Optional[bytes]:
    """
    Fetches qualification data from the database and generates an Excel report as a byte stream.

    Args:
        rfp_id: The RFP ID for which to generate the report.

    Returns:
        A bytes object containing the Excel file, or None if data is not found.
    """
    try:
        db_manager = DatabaseManager()

        print(f"🔍 [STEP 1] Fetching qualification results from database...")
        qualification_data = db_manager.get_qualification_results(rfp_id)

        if not qualification_data:
            print(f"❌ No qualification data found for RFP ID: {rfp_id}")
            return None

        print(f"✅ [STEP 1] Found qualification data")
        print(f"   Database row keys: {list(qualification_data.keys())}")

        # Extract the actual report from the JSONB column
        print(f"🔍 [STEP 2] Extracting qualification_report from JSONB...")
        report_data = qualification_data.get('qualification_report')
        if not report_data:
            print(f"❌ No 'qualification_report' field in database row")
            print(f"   Available keys: {list(qualification_data.keys())}")
            return None

        print(f"✅ [STEP 2] Extracted qualification_report JSONB")
        print(f"   Report data keys: {list(report_data.keys()) if isinstance(report_data, dict) else 'NOT A DICT'}")

        # Reconstruct Pydantic model from fetched data (lazy import)
        from agent.processors.new_rfp_qualifier import QualificationReport

        print(f"🔍 [STEP 3] Reconstructing QualificationReport Pydantic model...")
        try:
            report = QualificationReport(**report_data)
            print(f"✅ [STEP 3] Pydantic model created successfully")
        except Exception as pydantic_error:
            print(f"❌ [STEP 3] Failed to create Pydantic model: {pydantic_error}")
            import traceback
            traceback.print_exc()
            return None

        # Generate the Excel workbook (lazy import)
        from agent.processors.new_rfp_qualifier import create_qualification_excel_report

        print(f"🔍 [STEP 4] Generating Excel workbook...")
        try:
            workbook = create_qualification_excel_report(report)
            print(f"✅ [STEP 4] Excel workbook created")
        except Exception as excel_error:
            print(f"❌ [STEP 4] Failed to create Excel: {excel_error}")
            import traceback
            traceback.print_exc()
            return None

        # Save workbook to a byte stream
        print(f"🔍 [STEP 5] Saving to byte stream...")
        excel_buffer = io.BytesIO()
        workbook.save(excel_buffer)
        excel_buffer.seek(0)

        print(f"✅ [STEP 5] Excel bytes ready ({len(excel_buffer.getvalue())} bytes)")
        return excel_buffer.getvalue()

    except Exception as e:
        print(f"❌ UNEXPECTED ERROR in generate_and_download_qualification_excel: {e}")
        import traceback
        traceback.print_exc()
        return None

def generate_reasoning_excel(rfp_id: str) -> Optional[bytes]:
    """
    Generate a reasoning-only Excel report (detailed analysis with full reasoning).

    Args:
        rfp_id: The RFP ID for which to generate the reasoning report.

    Returns:
        A bytes object containing the Excel file, or None if data is not found.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, Border, Side, PatternFill

        db_manager = DatabaseManager()
        qualification_data = db_manager.get_qualification_results(rfp_id)

        if not qualification_data:
            print(f"❌ No qualification data found for reasoning report: {rfp_id}")
            return None

        report_data = qualification_data.get('qualification_report')
        if not report_data:
            return None

        from agent.processors.new_rfp_qualifier import QualificationReport
        report = QualificationReport(**report_data)

        # Create reasoning-focused workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Detailed Reasoning"

        # Header
        ws.cell(1, 1, f"QUALIFICATION REASONING REPORT - {report.client_and_opportunity}").font = Font(bold=True, size=14)
        ws.merge_cells('A1:E1')

        # Column headers
        headers = ["Criterion", "Score", "Reasoning", "Missing Data Justification", "Selected Option"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(3, col, header)
            cell.font = Font(bold=True, size=11)
            cell.fill = PatternFill(fill_type="solid", fgColor="4472C4")
            cell.font = Font(bold=True, size=11, color="FFFFFF")
            cell.alignment = Alignment(horizontal='center', wrap_text=True)

        # Data rows
        for i, analysis in enumerate(report.analyses, 4):
            ws.cell(i, 1, analysis.criterion)
            ws.cell(i, 2, f"{analysis.score}/4")
            ws.cell(i, 3, analysis.reasoning or "")
            ws.cell(i, 4, analysis.missing_data_justification or "")
            ws.cell(i, 5, analysis.selected_option)

            # Style
            for col in range(1, 6):
                cell = ws.cell(i, col)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.border = Border(
                    left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin')
                )

            # Dynamic row height
            reasoning_lines = len(analysis.reasoning.split('\n')) if analysis.reasoning else 0
            justification_lines = len(analysis.missing_data_justification.split('\n')) if analysis.missing_data_justification else 0
            ws.row_dimensions[i].height = max(60, max(reasoning_lines, justification_lines) * 15)

        # Column widths
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 70
        ws.column_dimensions['D'].width = 70
        ws.column_dimensions['E'].width = 50

        # Save to bytes
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        return excel_buffer.getvalue()

    except Exception as e:
        print(f"❌ Error generating reasoning Excel: {e}")
        import traceback
        traceback.print_exc()
        return None


def tool_requalify_rfp(rfp_id: str) -> Dict[str, Any]:
    """
    Re-run qualification for an existing RFP using stored document text.
    Useful when qualification data is missing but RFP document exists.

    Args:
        rfp_id: The RFP ID to re-qualify

    Returns:
        Dictionary with qualification results
    """
    try:
        print(f"\n🔄 [REQUALIFY] Re-qualifying RFP: {rfp_id}")

        # Get raw document text from database
        db_manager = DatabaseManager()
        raw_text = db_manager.get_raw_document_text(rfp_id)
        if not raw_text:
            return {
                'success': False,
                'error': 'No raw document text found',
                'message': f'Cannot re-qualify {rfp_id} - the original document text was not saved in the database. This RFP was likely processed before the raw text storage feature was added. To generate reports for this RFP, you need the user to re-upload the original document.'
            }

        print(f"✅ [REQUALIFY] Found {len(raw_text)} characters of document text")

        # Call qualification with the stored text
        return tool_qualify_rfp(context=raw_text, pdf_path=None)

    except Exception as e:
        import traceback
        return {
            'success': False,
            'error': str(e),
            'message': f'Failed to re-qualify: {str(e)}\n{traceback.format_exc()}'
        }


def generate_bid_plan_excel_bytes(rfp_id: str) -> Optional[bytes]:
    """
    Generate bid plan Excel and return as bytes (not file).

    Args:
        rfp_id: The RFP ID

    Returns:
        Bytes of the Excel file, or None if data missing
    """
    try:
        from openpyxl import load_workbook
        from agent.config.settings import settings
        import io

        complete_data = db.get_complete_rfp_data(rfp_id)

        deliverables = complete_data.get('deliverables')
        assignments = complete_data.get('assignments')
        raw_data = complete_data.get('raw_data')

        if not deliverables or not assignments:
            print(f"❌ Missing deliverables or assignments for {rfp_id}")
            return None

        # Load template from database
        print(f"📂 Loading bid plan template from database...")
        template_data = db.get_template('bid_plan_template')

        if not template_data:
            print(f"❌ Bid plan template not found in database")
            return None

        # Load template from bytes
        template_buffer = io.BytesIO(template_data)
        wb = load_workbook(template_buffer)

        # Use existing helper from excel_reports
        from .report_generators.excel_reports import _fill_deliverables_sheet, _fill_overview_sheet

        deliverables_sheet = None
        overview_sheet = None

        for name in wb.sheetnames:
            if "deliverable" in name.lower():
                deliverables_sheet = wb[name]
            elif "overview" in name.lower() or "team" in name.lower():
                overview_sheet = wb[name]

        if not deliverables_sheet:
            print(f"❌ Deliverables sheet not found in template")
            return None

        # Fill sheets
        _fill_deliverables_sheet(deliverables_sheet, deliverables, assignments)

        if overview_sheet and raw_data:
            _fill_overview_sheet(overview_sheet, raw_data['rfp_data'])

        # Save to bytes
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        return excel_buffer.getvalue()

    except Exception as e:
        print(f"❌ Error generating bid plan bytes: {e}")
        import traceback
        traceback.print_exc()
        return None


def generate_assignment_excel_bytes(rfp_id: str) -> Optional[bytes]:
    """
    Generate assignment report Excel and return as bytes (not file).

    Args:
        rfp_id: The RFP ID

    Returns:
        Bytes of the Excel file, or None if data missing
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
        from .report_generators.excel_reports import to_bullet_points
        import io

        assignments_data = db.get_rfp_assignments(rfp_id)

        if not assignments_data:
            print(f"❌ No assignment data found for {rfp_id}")
            return None

        report = assignments_data['assignment_report']

        wb = Workbook()
        ws = wb.active
        ws.title = "Assignment Analysis"

        # Headers
        headers = ["Deliverable Section", "Requirement", "Assigned Owner", "Reasoning", "Alternative Partners"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(1, col, header)
            cell.font = Font(name="Segoe UI", size=10, bold=True)
            cell.fill = PatternFill(fill_type="solid", fgColor="E7E6E6")
            cell.alignment = Alignment(horizontal='center', wrap_text=True)

        # Summary
        ws.cell(2, 1, f"Client: {report.get('client_and_opportunity', 'Unknown')}")
        ws.cell(3, 1, f"Analysis Date: {report.get('analysis_date', '')}")
        ws.cell(4, 1, f"Total: {report.get('total_deliverables', 0)}")
        ws.cell(5, 1, f"Granite: {report.get('granite_assigned', 0)}, Partners: {report.get('partner_assigned', 0)}")

        # Data rows
        start_row = 7
        for i, assignment in enumerate(report.get('assignments', []), start_row):
            ws.cell(i, 1, assignment.get('deliverable_section', ''))
            ws.cell(i, 2, to_bullet_points(assignment.get('deliverable_requirement', '')))
            ws.cell(i, 3, assignment.get('assigned_owner', ''))
            ws.cell(i, 4, to_bullet_points(assignment.get('reasoning', '')))

            alt_partners = assignment.get('alternative_partners', [])
            alt_text = "\n".join([f"• {partner}" for partner in alt_partners])
            ws.cell(i, 5, alt_text)

            for col in range(1, 6):
                cell = ws.cell(i, col)
                cell.font = Font(name="Segoe UI", size=10)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                cell.border = Border(
                    left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin')
                )

            ws.row_dimensions[i].height = max(30, len(assignment.get('deliverable_requirement', '').split('\n')) * 15)

        # Column widths
        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 60
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 45
        ws.column_dimensions['E'].width = 35

        # Save to bytes
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)

        return excel_buffer.getvalue()

    except Exception as e:
        print(f"❌ Error generating assignment bytes: {e}")
        import traceback
        traceback.print_exc()
        return None


def tool_download_qualification_report(rfp_id: str) -> Dict[str, Any]:
    """
    Generate qualification Excel reports (qualification + reasoning) and save to database.
    Files will appear in the downloads panel for all users.

    Args:
        rfp_id: The RFP ID

    Returns:
        Dictionary with success status and file count

    Example:
        tool_download_qualification_report("context_rfp_20251009_194115_af2fc196")
    """
    try:
        print(f"\n🔧 [DOWNLOAD TOOL] Generating reports for: {rfp_id}")

        excel_bytes = generate_and_download_qualification_excel(rfp_id)

        if not excel_bytes:
            print(f"❌ [DOWNLOAD TOOL] No qualification data found")
            return {
                'success': False,
                'error': 'No qualification data found',
                'message': f'No qualification data found for {rfp_id}. Run qualification first.'
            }

        print(f"✅ [DOWNLOAD TOOL] Generated {len(excel_bytes)} bytes")

        reasoning_bytes = generate_reasoning_excel(rfp_id)

        db_manager = DatabaseManager()
        saved_files = []

        qual_file_name = f'{rfp_id}_qualification_report.xlsx'
        file_id = db_manager.save_generated_file(
            session_id=None,
            file_name=qual_file_name,
            file_type='qualification',
            file_data=excel_bytes,
            rfp_id=rfp_id
        )
        saved_files.append(qual_file_name)
        print(f"✅ Saved {qual_file_name} (ID: {file_id})")

        if reasoning_bytes:
            reasoning_file_name = f'{rfp_id}_reasoning_report.xlsx'
            file_id = db_manager.save_generated_file(
                session_id=None,
                file_name=reasoning_file_name,
                file_type='reasoning',
                file_data=reasoning_bytes,
                rfp_id=rfp_id
            )
            saved_files.append(reasoning_file_name)
            print(f"✅ Saved {reasoning_file_name} (ID: {file_id})")
            message = f'✅ Generated 2 reports: {qual_file_name} and {reasoning_file_name}. Check downloads!'
        else:
            message = f'✅ Generated {qual_file_name}. Check downloads!'

        return {
            'success': True,
            'rfp_id': rfp_id,
            'files_saved': len(saved_files),
            'file_names': saved_files,
            'message': message
        }

    except Exception as e:
        import traceback
        print(f"❌ [DOWNLOAD TOOL] Error: {e}")
        return {
            'success': False,
            'error': str(e),
            'message': f'Failed to generate reports: {str(e)}'
        }


def tool_download_bid_plan_report(rfp_id: str) -> Dict[str, Any]:
    """
    Generate bid plan Excel reports (bid_plan + assignment) and save to database.
    Files will appear in the downloads panel for all users.

    Args:
        rfp_id: The RFP ID

    Returns:
        Dictionary with success status and file count

    Example:
        tool_download_bid_plan_report("context_rfp_20251009_194115_af2fc196")
    """
    try:
        print(f"\n🔧 [BID PLAN DOWNLOAD] Generating reports for: {rfp_id}")

        # Generate bid plan bytes
        bid_plan_bytes = generate_bid_plan_excel_bytes(rfp_id)

        if not bid_plan_bytes:
            print(f"❌ [BID PLAN DOWNLOAD] No bid plan data found")
            return {
                'success': False,
                'error': 'No bid plan data found',
                'message': f'No bid plan data found for {rfp_id}. Run bid planning first.'
            }

        print(f"✅ [BID PLAN DOWNLOAD] Generated {len(bid_plan_bytes)} bytes for bid plan")

        # Generate assignment bytes
        assignment_bytes = generate_assignment_excel_bytes(rfp_id)

        db_manager = DatabaseManager()
        saved_files = []

        # Save bid plan
        bid_plan_file_name = f'{rfp_id}_bid_plan.xlsx'
        file_id = db_manager.save_generated_file(
            session_id=None,
            file_name=bid_plan_file_name,
            file_type='bid_plan',
            file_data=bid_plan_bytes,
            rfp_id=rfp_id
        )
        saved_files.append(bid_plan_file_name)
        print(f"✅ Saved {bid_plan_file_name} (ID: {file_id})")

        # Save assignment if generated
        if assignment_bytes:
            assignment_file_name = f'{rfp_id}_assignment_report.xlsx'
            file_id = db_manager.save_generated_file(
                session_id=None,
                file_name=assignment_file_name,
                file_type='assignment',
                file_data=assignment_bytes,
                rfp_id=rfp_id
            )
            saved_files.append(assignment_file_name)
            print(f"✅ Saved {assignment_file_name} (ID: {file_id})")
            message = f'✅ Generated 2 reports: {bid_plan_file_name} and {assignment_file_name}. Check downloads!'
        else:
            message = f'✅ Generated {bid_plan_file_name}. Check downloads!'

        return {
            'success': True,
            'rfp_id': rfp_id,
            'files_saved': len(saved_files),
            'file_names': saved_files,
            'message': message
        }

    except Exception as e:
        print(f"❌ [BID PLAN DOWNLOAD] Error: {e}")
        return {
            'success': False,
            'error': str(e),
            'message': f'Failed to generate bid plan reports: {str(e)}'
        }


def tool_generate_report(
    rfp_id: str,
    report_type: str
) -> Dict[str, Any]:
    """
    Generate Excel reports on demand from stored data.
    Uses the LATEST data including all agent/user edits.

    Args:
        rfp_id: The RFP ID
        report_type: Type of report - "qualification", "bid_plan", "assignment", or "all"

    Returns:
        Dictionary with file paths to generated reports

    Examples:
        # Generate qualification report only
        tool_generate_report("KHDA_2024", "qualification")

        # Generate all reports at once
        tool_generate_report("KHDA_2024", "all")
    """
    try:
        from .report_generators.excel_reports import (
            generate_qualification_excel,
            generate_bid_plan_excel,
            generate_assignment_excel,
            generate_all_reports
        )

        if report_type == "qualification":
            file_path = generate_qualification_excel(rfp_id)
            return {
                'success': True,
                'rfp_id': rfp_id,
                'report_type': report_type,
                'file_path': file_path,
                'message': f'Generated qualification report: {file_path}'
            }

        elif report_type == "bid_plan":
            file_path = generate_bid_plan_excel(rfp_id)
            return {
                'success': True,
                'rfp_id': rfp_id,
                'report_type': report_type,
                'file_path': file_path,
                'message': f'Generated bid plan: {file_path}'
            }

        elif report_type == "assignment":
            file_path = generate_assignment_excel(rfp_id)
            return {
                'success': True,
                'rfp_id': rfp_id,
                'report_type': report_type,
                'file_path': file_path,
                'message': f'Generated assignment report: {file_path}'
            }

        elif report_type == "all":
            reports = generate_all_reports(rfp_id)
            return {
                'success': True,
                'rfp_id': rfp_id,
                'report_type': report_type,
                'reports': reports,
                'message': f'Generated all reports for {rfp_id}'
            }

        else:
            return {
                'success': False,
                'error': 'Invalid report_type',
                'message': f'report_type must be "qualification", "bid_plan", "assignment", or "all", got: {report_type}'
            }

    except Exception as e:
        return {
            'success': False,
            'error': str(e),
            'message': f'Failed to generate report: {str(e)}'
        }
