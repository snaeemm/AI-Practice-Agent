
from typing import Dict, Any, List, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
import io
from psycopg2.extras import Json
from pathlib import Path

from agent.database.db_manager import DatabaseManager
from agent.database.db_singleton import get_db

db = get_db()


def _get_logo_path() -> Path:
    """Get the path to the company logo file."""
    # Get the path relative to this file
    logo_path = Path(__file__).parent.parent.parent / "company_logo_vertical.png"
    return logo_path

def _is_title_slide(slide_data: Dict[str, Any]) -> bool:
    """Intelligently detect if a slide should be a title slide."""
    points = slide_data.get("points", [])

    # A title slide has no bullet points or very few (for subtitle)
    if len(points) == 0:
        return True
    elif len(points) == 1:
        # Check if it's a simple string (subtitle) or a nested bullet
        point = points[0]
        if isinstance(point, str) and len(point) < 100:
            return True

    return False


def _calculate_font_size(bullet_count: int) -> int:
    """Dynamically calculate font size based on bullet count."""
    if bullet_count <= 5:
        return 20  # Normal size
    elif bullet_count <= 8:
        return 18  # Slightly smaller
    else:
        return 16  # Compact


def _flatten_bullets(points: List[Any]) -> int:
    """Count total bullets including nested ones for font sizing."""
    count = 0
    for point in points:
        if isinstance(point, dict):
            count += 1  # Main point
            count += len(point.get("sub_points", []))  # Sub-points
        else:
            count += 1
    return count


def _create_presentation_bytes(slides: List[Dict[str, Any]]) -> bytes:
    """Create a PowerPoint presentation using proper layouts and bullet formatting."""
    prs = Presentation()

    # Set presentation dimensions to widescreen (16:9)
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)

    # Define Granite brand color scheme
    GRANITE_BLUE = RGBColor(41, 128, 185)      # Primary accent blue
    DARK_TEXT = RGBColor(44, 62, 80)           # Dark gray for text
    LIGHT_BG = RGBColor(245, 247, 250)         # Light background
    ACCENT_LIGHT = RGBColor(174, 214, 241)     # Light blue accent

    total_slides = len(slides)

    for idx, slide_data in enumerate(slides):
        title = slide_data.get("title", "Untitled")
        points = slide_data.get("points", [])
        is_title = idx == 0  # First slide is always a title slide

        if is_title:
            # Create a completely custom title slide with no template constraints
            slide_layout = prs.slide_layouts[6]  # Blank layout
            slide = prs.slides.add_slide(slide_layout)

            # Adaptive font size based on title length
            title_length = len(title)
            if title_length < 15:
                title_font_size = 48
            elif title_length < 25:
                title_font_size = 40
            elif title_length < 35:
                title_font_size = 34
            elif title_length < 50:
                title_font_size = 32
            else:
                title_font_size = 32  # Minimum 32pt for title slide

            # Position title with proper spacing from top
            title_top = 0.5  # Closer to top for title slides
            title_height_estimate = 1.2

            # Title text box - CENTERED
            title_box = slide.shapes.add_textbox(
                Inches(0.5),
                Inches(title_top),
                Inches(9),
                Inches(title_height_estimate)
            )
            title_frame = title_box.text_frame
            title_frame.word_wrap = True
            title_frame.vertical_anchor = MSO_ANCHOR.TOP

            p = title_frame.paragraphs[0]
            p.text = title
            p.font.size = Pt(title_font_size)
            p.font.bold = True
            p.font.color.rgb = GRANITE_BLUE
            p.alignment = PP_ALIGN.CENTER

            # Content/bullet points area - below title
            if points:
                content_top = title_top + title_height_estimate + 0.3
                content_height = 5.625 - content_top - 0.4  # Leave room for footer

                content_box = slide.shapes.add_textbox(
                    Inches(0.7),
                    Inches(content_top),
                    Inches(8.6),
                    Inches(content_height)
                )
                content_frame = content_box.text_frame
                content_frame.word_wrap = True
                content_frame.vertical_anchor = MSO_ANCHOR.TOP

                # Calculate font size based on bullet count
                total_bullets = _flatten_bullets(points)
                font_size = _calculate_font_size(total_bullets)

                for point_idx, point in enumerate(points):
                    if isinstance(point, dict):
                        # Main bullet with nested structure
                        if point_idx == 0:
                            cp = content_frame.paragraphs[0]
                        else:
                            cp = content_frame.add_paragraph()

                        cp.text = "• " + point.get("text", "")
                        cp.level = 0
                        cp.font.size = Pt(font_size)
                        cp.font.bold = True
                        cp.font.color.rgb = GRANITE_BLUE

                        # Add sub-points
                        sub_points = point.get("sub_points", [])
                        for sub_point in sub_points:
                            sub_p = content_frame.add_paragraph()
                            sub_p.text = "◦ " + sub_point
                            sub_p.level = 1
                            sub_p.font.size = Pt(font_size - 2)
                            sub_p.font.color.rgb = DARK_TEXT

                    else:
                        # Simple string bullet
                        if point_idx == 0:
                            cp = content_frame.paragraphs[0]
                        else:
                            cp = content_frame.add_paragraph()

                        cp.text = "• " + point
                        cp.level = 0
                        cp.font.size = Pt(font_size)
                        cp.font.color.rgb = DARK_TEXT

            # Add logo watermark at mid right with background
            logo_path = _get_logo_path()
            if logo_path.exists():
                try:
                    # Add a subtle dark background rectangle behind logo
                    bg_shape = slide.shapes.add_shape(
                        1,  # Rectangle shape type
                        Inches(9.505),  # Slightly left of logo
                        Inches(2.4),  # Mid-right vertical position
                        Inches(.4),  # Width to accommodate logo
                        Inches(0.9)   # Height to accommodate logo
                    )
                    # Style the background
                    bg_shape.fill.solid()
                    bg_shape.fill.fore_color.rgb = RGBColor(44, 62, 80)  # Dark gray/blue
                    bg_shape.line.fill.background()  # No border

                    # Position logo on top of background
                    logo_pic = slide.shapes.add_picture(
                        str(logo_path),
                        Inches(9.6),  # Right side
                        Inches(2.5),  # Mid-right vertical position
                        height=Inches(0.7)  # Small watermark size
                    )
                except Exception as e:
                    pass  # Silently fail if logo can't be added

        else:
            # Use title and content layout
            slide_layout = prs.slide_layouts[1]  # Title and content layout
            slide = prs.slides.add_slide(slide_layout)

            # Set title
            title_shape = slide.shapes.title
            title_shape.text = title
            title_frame = title_shape.text_frame
            title_frame.word_wrap = True

            for paragraph in title_frame.paragraphs:
                paragraph.font.size = Pt(44)
                paragraph.font.bold = True
                paragraph.font.color.rgb = GRANITE_BLUE
                paragraph.alignment = PP_ALIGN.LEFT

            # Set content (bullets)
            if len(slide.placeholders) > 1 and points:
                body_shape = slide.placeholders[1]
                text_frame = body_shape.text_frame
                text_frame.clear()
                text_frame.word_wrap = True

                # Calculate font size based on bullet count
                total_bullets = _flatten_bullets(points)
                font_size = _calculate_font_size(total_bullets)

                for point_idx, point in enumerate(points):
                    if isinstance(point, dict):
                        # Main bullet with nested structure
                        if point_idx == 0:
                            p = text_frame.paragraphs[0]
                        else:
                            p = text_frame.add_paragraph()

                        p.text = point.get("text", "")
                        p.level = 0
                        p.font.size = Pt(font_size)
                        p.font.bold = True
                        p.font.color.rgb = GRANITE_BLUE

                        # Add sub-points
                        sub_points = point.get("sub_points", [])
                        for sub_point in sub_points:
                            sub_p = text_frame.add_paragraph()
                            sub_p.text = sub_point
                            sub_p.level = 1
                            sub_p.font.size = Pt(font_size - 2)
                            sub_p.font.color.rgb = DARK_TEXT

                    else:
                        # Simple string bullet
                        if point_idx == 0:
                            p = text_frame.paragraphs[0]
                        else:
                            p = text_frame.add_paragraph()

                        p.text = point
                        p.level = 0
                        p.font.size = Pt(font_size)
                        p.font.color.rgb = DARK_TEXT

            # Add logo watermark at mid right with background
            logo_path = _get_logo_path()
            if logo_path.exists():
                try:
                    # Add a subtle dark background rectangle behind logo
                    bg_shape = slide.shapes.add_shape(
                        1,  # Rectangle shape type
                        Inches(9.505),  # Slightly left of logo
                        Inches(2.4),  # Mid-right vertical position
                        Inches(.4),  # Width to accommodate logo
                        Inches(0.9)   # Height to accommodate logo
                    )
                    # Style the background
                    bg_shape.fill.solid()
                    bg_shape.fill.fore_color.rgb = RGBColor(44, 62, 80)  # Dark gray/blue
                    bg_shape.line.fill.background()  # No border

                    # Position logo on top of background
                    logo_pic = slide.shapes.add_picture(
                        str(logo_path),
                        Inches(9.6),  # Right side
                        Inches(2.5),  # Mid-right vertical position
                        height=Inches(0.7)  # Small watermark size
                    )
                except Exception as e:
                    pass  # Silently fail if logo can't be added

            # Add footer with slide number
            if len(slide.placeholders) > 2:
                # Try to use footer placeholder if available
                try:
                    footer = slide.placeholders[2]
                    footer.text = f"Slide {idx}/{total_slides}"
                except:
                    pass

    # Save to bytes
    f = io.BytesIO()
    prs.save(f)
    f.seek(0)
    return f.read()

def _validate_presentation_structure(slides: List[Dict[str, Any]]) -> tuple[bool, str]:
    """
    Validate presentation structure and provide feedback.

    Returns:
        Tuple of (is_valid, message)
    """
    warnings = []

    for idx, slide in enumerate(slides):
        points = slide.get("points", [])
        total_bullets = _flatten_bullets(points)

        # Check for too many bullets
        if total_bullets > 8:
            warnings.append(f"Slide {idx + 1} has {total_bullets} bullet points. Consider splitting into 2 slides (ideal: 3-7 bullets)")

        # Validate structure format
        for point_idx, point in enumerate(points):
            if isinstance(point, dict):
                # Check required fields
                if "text" not in point:
                    return False, f"Slide {idx + 1}, bullet {point_idx + 1}: Nested bullet missing 'text' field"

                # Check sub-points format
                sub_points = point.get("sub_points", [])
                if sub_points and not isinstance(sub_points, list):
                    return False, f"Slide {idx + 1}, bullet {point_idx + 1}: 'sub_points' must be a list"

            elif not isinstance(point, str):
                return False, f"Slide {idx + 1}, bullet {point_idx + 1}: Point must be string or nested object with 'text' field"

    # Return validation result with warnings
    if warnings:
        warning_text = "\n".join(["⚠️ " + w for w in warnings])
        return True, f"Presentation is valid but has suggestions:\n{warning_text}\nProceeding with save..."
    else:
        return True, "✅ Presentation structure is optimal!"


def tool_save_presentation_structure(
    presentation_title: str,
    slides: List[Dict[str, Any]],
    presentation_description: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Saves the structured presentation outline to the database.

    Args:
        presentation_title: The title of the presentation (must be unique).
        slides: A list of slides, where each slide is a dictionary with "title" and "points".
        presentation_description: Optional description of the presentation.

    Returns:
        A dictionary with the success status.
    """
    try:
        # Validate structure
        is_valid, validation_msg = _validate_presentation_structure(slides)
        if not is_valid:
            return {
                "success": False,
                "error": "Invalid presentation structure",
                "message": validation_msg
            }

        # Check if presentation already exists
        existing = db.get_presentation_by_title(presentation_title)

        structure = {"slides": slides}

        if existing:
            # Update existing presentation
            with db._get_connection() as conn:
                with conn.cursor() as cursor:
                    cursor.execute("""
                        UPDATE presentations
                        SET presentation_structure = %s,
                            presentation_description = %s
                        WHERE presentation_title = %s
                    """, (Json(structure), presentation_description, presentation_title))
                    conn.commit()
            return {
                "success": True,
                "message": f"Presentation '{presentation_title}' updated successfully with {len(slides)} slides\n{validation_msg}"
            }
        else:
            # Create new presentation
            with db._get_connection() as conn:
                with conn.cursor() as cursor:
                    cursor.execute("""
                        INSERT INTO presentations
                        (presentation_title, presentation_description, presentation_structure)
                        VALUES (%s, %s, %s)
                        RETURNING id
                    """, (presentation_title, presentation_description, Json(structure)))
                    conn.commit()
                    pres_id = cursor.fetchone()[0]
            return {
                "success": True,
                "message": f"Presentation '{presentation_title}' saved successfully with {len(slides)} slides\n{validation_msg}"
            }
    except Exception as e:
        return {
            "success": False,
            "error": str(e),
            "message": f"Failed to save presentation structure: {str(e)}"
        }

def tool_generate_presentation_from_db(
    rfp_id: str,
    file_name: str
) -> Dict[str, Any]:
    """
    Generates a PowerPoint presentation from a structured outline stored in the database.

    Args:
        rfp_id: The ID of the RFP associated with the presentation.
        file_name: The desired name for the generated PowerPoint file.

    Returns:
        A dictionary with the success status and the presentation content as bytes.
    """
    try:
        presentation_data = db.get_presentation_structure(rfp_id)
        if not presentation_data or "slides" not in presentation_data:
            return {"success": False, "error": "No presentation structure found", "message": f"No presentation structure found for RFP: {rfp_id}"}

        slides = presentation_data["slides"]
        ppt_bytes = _create_presentation_bytes(slides)

        return {
            "success": True,
            "file_name": file_name,
            "content": ppt_bytes,
            "message": f"Successfully generated presentation {file_name} from DB for RFP: {rfp_id}"
        }
    except Exception as e:
        return {
            "success": False,
            "error": str(e),
            "message": f"Failed to generate presentation from DB: {str(e)}"
        }
